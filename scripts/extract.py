"""Extract PDF/DOCX/TXT from stdin. No shell commands or macros are executed."""
import sys, json, base64, io

try:
    item = json.load(sys.stdin)
    raw = base64.b64decode(item['base64'], validate=True)
    kind = item['extension']
    if kind == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError('PDF có mật khẩu. Vui lòng tải bản không có mật khẩu.')
        if len(reader.pages) > 300:
            raise ValueError('Tối đa 300 trang trong một tài liệu.')
        text = '\n\n'.join(f'[Trang {i+1}]\n{page.extract_text() or ""}' for i, page in enumerate(reader.pages))
        if len(''.join(page.extract_text() or '' for page in reader.pages).strip()) < 30:
            raise ValueError('PDF quét chưa có lớp chữ. Hãy OCR và kiểm tra tài liệu trước khi tải lại.')
    elif kind == 'docx':
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Tệp Word giải nén vượt giới hạn 50 MB.')
        from docx import Document
        doc = Document(io.BytesIO(raw))
        text = '\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
    else:
        text = raw.decode('utf-8-sig')
    if len(text) > 300000:
        raise ValueError('Nội dung vượt 300.000 ký tự. Hãy chia nhỏ tài liệu.')
    print(json.dumps({'text': text}, ensure_ascii=False))
except Exception as exc:
    print(json.dumps({'error': str(exc)}, ensure_ascii=False))
    sys.exit(1)
